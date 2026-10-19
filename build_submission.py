import os
import shutil
import zipfile
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import build_war

def zip_directory_contents(src_dir, target_jar):
    print(f"Creating {target_jar} from {src_dir}...")
    with zipfile.ZipFile(target_jar, 'w', zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(src_dir):
            for file in files:
                abs_path = os.path.join(root, file)
                rel_path = os.path.relpath(abs_path, src_dir)
                zf.write(abs_path, rel_path)

def generate_readme(file_path):
    print(f"Generating {file_path}...")
    doc = docx.Document()

    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Guess Market - Exercise 3 README (Client-Server)")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Bonus Highlight Banner (At the very beginning as required by instructions)
    p_bonus_banner = doc.add_paragraph()
    run_b1 = p_bonus_banner.add_run("★ Bonus Feature Implemented: Real-Time Multi-User Chat Room (5 Points Bonus) ★\n")
    run_b1.font.name = 'Arial'
    run_b1.font.bold = True
    run_b1.font.size = Pt(13)
    run_b1.font.color.rgb = RGBColor(0, 102, 204)
    run_b2 = p_bonus_banner.add_run("A live, multi-user global chat room is integrated into both the Tomcat backend (/api/chat) and the JavaFX GUI client (under the '💬 Chat Room' tab), enabling all active market participants to broadcast and receive messages in real time.")
    run_b2.font.name = 'Arial'
    run_b2.font.size = Pt(11)

    # Section 1: Student Information
    h1 = doc.add_heading("1. Student & Contact Information", level=1)
    h1.style.font.name = 'Arial'
    h1.style.font.color.rgb = RGBColor(0, 51, 102)
    p1 = doc.add_paragraph()
    p1.add_run(" • Name: Rami Zargian\n • ID: 324081231\n • Email: ramizargian@gmail.com").font.name = 'Arial'

    # Section 2: GitHub Repository
    h2 = doc.add_heading("2. GitHub Project Repository", level=1)
    h2.style.font.name = 'Arial'
    h2.style.font.color.rgb = RGBColor(0, 51, 102)
    p2 = doc.add_paragraph()
    p2.add_run(" • Repository URL: https://github.com/rami968/Guess-Market.git").font.name = 'Arial'

    # Section 3: Architecture & Design Choices
    h3 = doc.add_heading("3. System Architecture & Design Choices", level=1)
    h3.style.font.name = 'Arial'
    h3.style.font.color.rgb = RGBColor(0, 51, 102)
    p3 = doc.add_paragraph()
    p3.add_run(
        "Guess Market (Exercise 3) transitions the prediction market from a local desktop application into a distributed multi-tier Client-Server application.\n\n"
        "Architecture Overview:\n"
        " • Server Tier (Tomcat 10.1 / Jakarta EE): Packaged as a single standalone WAR (guess_market.war). Houses the centralized EngineService instance, in-memory state, and RESTful Servlets handling JSON requests and multipart XML streams.\n"
        " • Client Tier (JavaFX 21+): Rich client application connecting over HTTP. Communicates strictly with the server through HttpClientManager. Features periodic background pulling (every 1 second) to synchronize market state and chat messages automatically.\n"
        " • Core Engine & DTOs: Houses immutable Data Transfer Objects, prediction market math, Order Book matching engine with Auto-Minting, and event accounting.\n\n"
        "Key Design Choices & Adherence to Guidelines:\n"
        " 1. Pure In-Memory Server State: Conforming strictly with the requirement that the examiner server has no disk write permissions, XML files are parsed on-the-fly directly from input streams without saving files to disk. Furthermore, all state resides in memory across servlet invocations, meaning if the server restarts, all state resets cleanly.\n"
        " 2. Dynamic User Registration: Users register upon launching the client through a dedicated Login screen, ensuring username uniqueness. No hardcoded users exist in XML files.\n"
        " 3. Market Maker Dynamic Attribution: Any connected trader can upload an XML configuration file. The uploading user automatically becomes the designated Market Maker for all newly added events.\n"
        " 4. Event Accumulation: New XML uploads append events to the system instead of replacing existing ones.\n"
        " 5. Efficient Pull Synchronization: The server maintains an atomic state version counter. The client lightweightly queries the version every 1 second and only pulls heavy event/user datasets when state has changed. Chat messages are retrieved using delta-fetching based on message index.\n"
        " 6. Event Closure Balance Zeroing: Event closure distributes gross payouts and commission to participants, guaranteeing the event account balance drops exactly to $0.00.\n"
        " 7. Responsive GUI Layout: Built using SplitPanes, ScrollPanes, and FlowPanes with minimum stage bounds (1100x700), ensuring all controls, tables, and Order Book statistics (LAST, BID, ASK, MID, SPREAD) remain fully visible and scale gracefully upon resizing."
    ).font.name = 'Arial'

    # Section 4: Main Classes Description
    h4 = doc.add_heading("4. Main Classes Description", level=1)
    h4.style.font.name = 'Arial'
    h4.style.font.color.rgb = RGBColor(0, 51, 102)
    p4 = doc.add_paragraph()
    p4.add_run(
        "Server Module:\n"
        " • server.context.EngineContext: Application singleton holder managing the thread-safe EngineService, shared Gson serializer, and global chat message queue.\n"
        " • server.utils.ServletUtils: Standardized utility for reading JSON requests, extracting parameters, and writing structured responses.\n"
        " • server.servlets.AuthServlet (/api/login): Handles user registration and validation, rejecting duplicates with HTTP 409 Conflict.\n"
        " • server.servlets.UsersServlet (/api/users): Serves personal user profiles or other users overview summaries for the table.\n"
        " • server.servlets.ChargeServlet (/api/user/charge): Processes user balance deposits.\n"
        " • server.servlets.UploadEventServlet (/api/events/upload): Accepts multipart XML uploads and streams them directly into the engine in-memory without touching disk.\n"
        " • server.servlets.EventsServlet (/api/events): Lists loaded events and delivers real-time Order Book and LMSR status.\n"
        " • server.servlets.EventActionServlet (/api/events/action): Dispatches MM activation, Order Book limit orders, LMSR share purchases, and event closures.\n"
        " • server.servlets.SyncServlet (/api/sync): Lightweight endpoint returning current state version for client pull checks.\n"
        " • server.servlets.ChatServlet (/api/chat): Bonus chat broadcast and delta-fetching endpoint.\n\n"
        "Client (UI) Module:\n"
        " • ui.http.HttpClientManager: Centralized HTTP client using Java standard java.net.http.HttpClient for all backend communication.\n"
        " • ui.controller.LoginController: Manages connection to Tomcat and initial user registration.\n"
        " • ui.controller.MainController: Coordinates layout, user balance display, deposit/upload dialogs, and background pull executor.\n"
        " • ui.controller.EventsTabController: Handles event filtering via Toggle Buttons, drill-down details, Order Book BID/ASK tables, and trade submission.\n"
        " • ui.controller.UsersTabController: Displays personal holdings, full transaction history log, and table of other participants.\n"
        " • ui.controller.ChatTabController: Manages live group chat room interaction."
    ).font.name = 'Arial'

    # Section 5: Execution Instructions
    h5 = doc.add_heading("5. Execution Instructions", level=1)
    h5.style.font.name = 'Arial'
    h5.style.font.color.rgb = RGBColor(0, 51, 102)
    p5 = doc.add_paragraph()
    p5.add_run(
        "Step 1: Start Server (Apache Tomcat 10.1)\n"
        " 1. Copy 'guess_market.war' into your Tomcat 'webapps' directory (e.g. C:\\Tomcat\\webapps).\n"
        " 2. Launch Tomcat by executing 'startup.bat' located in Tomcat's 'bin' directory.\n"
        " 3. Verify server is running at: http://localhost:8080/guess_market/ping\n\n"
        "Step 2: Launch Client Application\n"
        " 1. Open the 'Client' folder included in the submission package.\n"
        " 2. Double-click 'run.bat'.\n"
        " 3. Enter your username on the Login screen and click 'Connect & Enter Market'.\n"
        " 4. Click 'Upload XML Events' to add new prediction events, or 'Charge Balance' to deposit funds."
    ).font.name = 'Arial'

    doc.save(file_path)
    print(f"Successfully generated {file_path}")

def main():
    print("=== Step 1: Building WAR ===")
    build_war.build_war()

    print("=== Step 2: Compiling UI ===")
    os.makedirs("out/production/UI", exist_ok=True)
    ui_files = build_war.find_java_files("UI/src")
    build_war.run_cmd(f'javac --module-path "lib/javafx-sdk/lib" --add-modules javafx.controls,javafx.fxml -cp "lib/*;out/production/Engine" -d out/production/UI {" ".join(ui_files)}')
    if os.path.exists("UI/src/ui/resources"):
        shutil.copytree("UI/src/ui/resources", "out/production/UI/ui/resources", dirs_exist_ok=True)

    print("=== Step 3: Preparing Submission Package ===")
    submission_root = "Rami_Zargian_324081231_EX3"
    client_dir = os.path.join(submission_root, "Client")
    if os.path.exists(submission_root):
        shutil.rmtree(submission_root)
    os.makedirs(client_dir, exist_ok=True)

    # 1. Copy WAR to root of submission
    war_src = "out/artifacts/Server_war/guess_market.war"
    shutil.copy2(war_src, os.path.join(submission_root, "guess_market.war"))

    # 2. Package Client JARs into Client/
    zip_directory_contents("out/production/UI", os.path.join(client_dir, "UI.jar"))
    zip_directory_contents("out/production/Engine", os.path.join(client_dir, "Engine.jar"))

    # 3. Copy runtime lib folder to Client/lib
    target_lib = os.path.join(client_dir, "lib")
    shutil.copytree("lib", target_lib, dirs_exist_ok=True)

    # 4. Create standalone run.bat into Client/
    client_bat = os.path.join(client_dir, "run.bat")
    with open(client_bat, "w", encoding="utf-8") as f:
        f.write('@echo off\n')
        f.write('echo Starting Guess Market JavaFX Client...\n')
        f.write('echo Note: Please ensure Apache Tomcat is running on port 8080 with guess_market.war deployed.\n')
        f.write('java --enable-native-access=javafx.graphics --module-path "%~dp0lib\\javafx-sdk\\lib" --add-modules javafx.controls,javafx.fxml -cp "%~dp0UI.jar;%~dp0Engine.jar;%~dp0lib\\*" ui.MainApp\n')
        f.write('pause\n')

    # 5. Generate README.docx
    readme_path = os.path.join(submission_root, "README.docx")
    generate_readme(readme_path)
    shutil.copy2(readme_path, "README.docx")

    # 6. Create final ZIP
    zip_name = "Rami_Zargian_324081231_EX3.zip"
    if os.path.exists(zip_name):
        os.remove(zip_name)

    print(f"Creating submission zip: {zip_name}...")
    with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(submission_root):
            for file in files:
                if file.startswith("~$"):
                    continue
                abs_path = os.path.join(root, file)
                rel_path = os.path.relpath(abs_path, submission_root)
                zf.write(abs_path, rel_path)

    print(f"\nSUCCESS! Submission package created: {zip_name}")

if __name__ == "__main__":
    main()
